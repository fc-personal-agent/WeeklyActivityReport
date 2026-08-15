"""Build a Word document from a JSON transcription of a handwritten PDF.

Usage:
    python build_docx.py <input.json> <output.docx>

Input JSON shape:
{
  "title": "Transcription of <source filename>",
  "source_pdf": "path/to/original.pdf",
  "pages": [
    {"page_number": 1, "paragraphs": ["line one", "line two", ...]},
    ...
  ],
  "notes": ["optional list of flagged unclear/illegible items"]
}
"""
import json
import sys

from docx import Document
from docx.shared import Pt


def build(input_json: str, output_path: str) -> None:
    with open(input_json, "r", encoding="utf-8-sig") as f:
        data = json.load(f)

    doc = Document()

    title = data.get("title", "Handwritten Document Transcription")
    doc.add_heading(title, level=0)

    source_pdf = data.get("source_pdf")
    if source_pdf:
        meta = doc.add_paragraph()
        run = meta.add_run(f"Source: {source_pdf}")
        run.italic = True
        run.font.size = Pt(9)

    for page in data.get("pages", []):
        page_number = page.get("page_number")
        doc.add_heading(f"Page {page_number}", level=1)
        paragraphs = page.get("paragraphs", [])
        if not paragraphs:
            doc.add_paragraph("[no legible text found on this page]")
        for para_text in paragraphs:
            doc.add_paragraph(para_text)

    notes = data.get("notes", [])
    if notes:
        doc.add_heading("Notes on Unclear or Illegible Text", level=1)
        for note in notes:
            doc.add_paragraph(note, style="List Bullet")

    doc.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_docx.py <input.json> <output.docx>")
        sys.exit(1)
    build(sys.argv[1], sys.argv[2])
