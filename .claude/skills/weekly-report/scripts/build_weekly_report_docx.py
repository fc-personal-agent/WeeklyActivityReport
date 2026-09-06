"""Build the final "Weekly SRE Update" Word document from a JSON summary of
a handwritten PDF (or of a prior transcription / executive summary).

Usage:
    python build_weekly_report_docx.py <input.json> <output.docx>

Input JSON shape:
{
  "title": "Weekly SRE Update - AssetMark - August 24-28, 2026",
  "source": "path/to/original.pdf or transcription.docx",
  "summary": "Brief overview paragraph.",
  "key_findings": ["...", "..."],
  "follow_up_items": ["...", "..."],
  "additional_notes": ["...", "..."]
}
"""
import json
import sys

from docx import Document
from docx.shared import Pt

# Number of blank "List Bullet" lines to append after real content, so the
# user has ready-made bullet slots to type more into directly in Word.
TRAILING_BLANK_BULLETS = {
    "Follow-up Items / Open Questions": 2,
    "Additional Notes": 8,
}


def add_bullets(doc: Document, heading: str, items: list[str], trailing_blanks: int = 0) -> None:
    if not items and not trailing_blanks:
        return
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
    for _ in range(trailing_blanks):
        doc.add_paragraph("", style="List Bullet")


def build(input_json: str, output_path: str) -> None:
    with open(input_json, "r", encoding="utf-8-sig") as f:
        data = json.load(f)

    doc = Document()

    title = data.get("title", "Weekly SRE Update")
    doc.add_heading(title, level=0)

    source = data.get("source")
    if source:
        meta = doc.add_paragraph()
        run = meta.add_run(f"Source: {source}")
        run.italic = True
        run.font.size = Pt(9)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(data.get("summary", "").strip() or "No summary provided.")

    add_bullets(doc, "Key Findings & Actions", data.get("key_findings", []))
    add_bullets(
        doc,
        "Follow-up Items / Open Questions",
        data.get("follow_up_items", []),
        TRAILING_BLANK_BULLETS["Follow-up Items / Open Questions"],
    )
    add_bullets(
        doc,
        "Additional Notes",
        data.get("additional_notes", []),
        TRAILING_BLANK_BULLETS["Additional Notes"],
    )

    doc.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_weekly_report_docx.py <input.json> <output.docx>")
        sys.exit(1)
    build(sys.argv[1], sys.argv[2])
