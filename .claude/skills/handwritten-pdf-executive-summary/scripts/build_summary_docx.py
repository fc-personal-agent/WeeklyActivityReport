"""Build an executive-summary Word document from a JSON summary of a
handwritten PDF (or of a prior verbatim transcription).

Usage:
    python build_summary_docx.py <input.json> <output.docx>

Input JSON shape:
{
  "title": "Executive Summary of <source filename>",
  "source": "path/to/original.pdf or transcription.docx",
  "summary": "Brief overview paragraph.",
  "key_findings": ["...", "..."],
  "follow_up_items": ["...", "..."],
  "confirmation_needed": ["...", "..."]
}
"""
import json
import sys

from docx import Document
from docx.shared import Pt


def add_bullets(doc: Document, heading: str, items: list[str]) -> None:
    if not items:
        return
    doc.add_heading(heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build(input_json: str, output_path: str) -> None:
    with open(input_json, "r", encoding="utf-8-sig") as f:
        data = json.load(f)

    doc = Document()

    title = data.get("title", "Executive Summary")
    doc.add_heading(title, level=0)

    source = data.get("source")
    if source:
        meta = doc.add_paragraph()
        run = meta.add_run(f"Source: {source}")
        run.italic = True
        run.font.size = Pt(9)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(data.get("summary", "").strip() or "No summary provided.")

    add_bullets(doc, "Key Findings & Actions", data.get("key_findings", []))
    add_bullets(doc, "Follow-up Items / Open Questions", data.get("follow_up_items", []))
    add_bullets(
        doc,
        "Areas Requiring Human Confirmation (Unclear Handwriting)",
        data.get("confirmation_needed", []),
    )

    doc.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_summary_docx.py <input.json> <output.docx>")
        sys.exit(1)
    build(sys.argv[1], sys.argv[2])
